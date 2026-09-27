import pandas as pd
from docx import Document

df = pd.read_csv("run_headless_results.csv")
# Remove the two corrupted rows
df = df[~((df['Model'] == 'diffusiongemma-26b-a4b-it') & (df['Scenario'].isin(['Criminal Law: Trespass', 'Tort Law: Negligence (Slip and Fall)'])))]
df.to_csv("run_headless_results.csv", index=False)

# Rebuild markdown without those two sections
with open("supplementary_evidence.md", "r") as f:
    full_md = f.read()

sections = full_md.split("## Model: ")
final_md = sections[0]
for sec in sections[1:]:
    header = sec.split("\n")[0]
    if "diffusiongemma-26b-a4b-it" in header and ("Criminal Law: Trespass" in header or "Tort Law: Negligence (Slip and Fall)" in header):
        continue
    final_md += "## Model: " + sec

with open("supplementary_evidence.md", "w") as f:
    f.write(final_md)

# Rebuild DOCX
doc = Document()
doc.add_heading("Supplementary Evidence: Full Evaluation Transcripts", 0)
for line in final_md.split("\n"):
    if line.startswith("# "): doc.add_heading(line[2:], level=1)
    elif line.startswith("## "): doc.add_heading(line[3:], level=2)
    elif line.startswith("### "): doc.add_heading(line[4:], level=3)
    elif line.startswith("**") and "**" in line[2:]:
        p = doc.add_paragraph()
        p.add_run(line).bold = True
    elif line.startswith("```"): continue
    else:
        if line.strip(): doc.add_paragraph(line)

doc.save("Supplementary_Evidence.docx")
print("Cleaned!")
